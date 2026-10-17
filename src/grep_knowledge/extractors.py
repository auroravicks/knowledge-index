from __future__ import annotations

import json
from pathlib import Path

import nbformat
from docx import Document
from pypdf import PdfReader
from pptx import Presentation


def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf(path)
    if suffix == ".docx":
        return _docx(path)
    if suffix == ".pptx":
        return _pptx(path)
    if suffix == ".ipynb":
        return _ipynb(path)
    if suffix in {".txt", ".md", ".csv", ".tsv", ".py", ".r"}:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported extension: {suffix}")


def _pdf(path: Path) -> str:
    output = []
    for number, page in enumerate(PdfReader(path).pages, 1):
        output += [f"\n[[PAGE {number}]]\n", page.extract_text() or ""]
    return "\n".join(output)


def _docx(path: Path) -> str:
    return "\n".join(
        f"[[PARAGRAPH {number}]] {paragraph.text}"
        for number, paragraph in enumerate(Document(path).paragraphs, 1)
        if paragraph.text.strip()
    )


def _pptx(path: Path) -> str:
    output = []
    for number, slide in enumerate(Presentation(path).slides, 1):
        text = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        output += [f"\n[[SLIDE {number}]]\n", "\n".join(text)]
    return "\n".join(output)


def _ipynb(path: Path) -> str:
    notebook = nbformat.read(path, as_version=4)
    output = []
    for number, cell in enumerate(notebook.cells, 1):
        output += [f"\n[[CELL {number} — {cell.cell_type.upper()}]]\n", cell.source]
    return "\n".join(output)

